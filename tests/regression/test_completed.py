"""completed_on -- the job is done, dated, and a re-import leaves it alone.

"Complete" means the work is done: shipped or installed. It is NOT payment
(collection status stays its own field), and it does NOT clear the Live
Tracker bucket: tracker_status stays as provenance, and the rule is that a
project with completed_on set is not live, whatever its bucket.

Five decisions are asserted, each with an obvious wrong version:

 1. STORED AS GIVEN, like next_action_on. null clears it ("Reopen").
 2. AN UNREADABLE DATE IS REFUSED, naming the value; a date after today is
    refused, naming both dates -- through the frozen clock, never the wall
    clock.
 3. A DATE BEFORE THE DEAL DATE IS ALLOWED WITH A WARNING naming both dates.
    A messy deal date must never block marking a real job done.
 4. IT SURVIVES: a fresh process reads it back off disk and out of the
    changelog, and a re-import through normalize.run keeps it while the
    importer re-sets tracker_status from the sheet's colour.
 5. TWO CUSTOMERS HOLDING ONE NUMBER are completed independently, and a
    re-import leaves both records exactly as they were (merge.py treats a
    shared number as ambiguous and touches neither -- a known limit, not
    changed here).

Driven through server.mcp.call_tool and pipeline/normalize.run. Fixture names
are invented.
"""
import json
import shutil
import subprocess
import sys
import tempfile
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from lib.harness import Result, Store, company, project  # noqa: E402

TODAY = date(2026, 8, 9)
YELLOW = "FFFFFF00"     # the fixture legend's second bucket -> action_owner


def _rec(s, pno, cid="acme"):
    return next((p for p in s.read("projects") or []
                 if str(p.get("project_no")) == pno and p.get("company_id") == cid), None)


def _f(s, pno, key, cid="acme"):
    p = _rec(s, pno, cid)
    return p.get(key) if p else None


def run(server, crm_dir=None):
    r = Result("completed", since="0.1.43")
    crm = Path(crm_dir) if crm_dir else \
        Path(__file__).resolve().parents[2] / "plugins/unrivaled-solutions/skills/crm"
    srv = server
    if srv is None:
        from lib.harness import load_server
        srv = load_server(str(crm))
    if not r.check("server exposes _today()", callable(getattr(srv, "_today", None))):
        return r
    srv._today = lambda: TODAY
    r.check("completed_on is a project field",
            "completed_on" in getattr(srv, "PROJECT_FIELDS", set()))
    s = Store(srv)

    def seed():
        s.reset(companies=[company("acme", "Ace Manufacturing"),
                           company("beta", "Beta Works")],
                projects=[project("4521", "acme", date="2026-03-14 00:00:00",
                                  tracker_status="action_admin",
                                  collection_status="open"),
                          project("4522", "acme", date="3/14/2026"),
                          project("4523", "acme", date="not a date"),
                          project("4524", "acme")])

    # ---- 1. set, stored as given; clear --------------------------------------
    r.section("set and clear")
    seed()
    res = s.call("update_project", project_no="4521", company_id="acme",
                 fields={"completed_on": "2026-08-01"})
    r.check("update_project accepts completed_on",
            res.get("ok") is True and res.get("project", {}).get("completed_on") == "2026-08-01",
            json.dumps(res)[:200])
    r.check("it is on disk exactly as given", _f(s, "4521", "completed_on") == "2026-08-01")
    r.check("marking complete does NOT clear the Live Tracker bucket",
            _f(s, "4521", "tracker_status") == "action_admin")
    r.check("nor touch the collection status -- complete is not paid",
            _f(s, "4521", "collection_status") == "open")
    r.check("a clean completion carries no warning", not res.get("warnings"),
            json.dumps(res.get("warnings")))
    res = s.call("get_project", project_no="4521", company_id="acme")
    r.check("get_project reads it back",
            res.get("ok") and res.get("project", {}).get("completed_on") == "2026-08-01")
    res = s.call("update_project", project_no="4522", company_id="acme",
                 fields={"completed_on": "8/1/26"})
    r.check("a tracker-style date is kept verbatim",
            res.get("ok") is True and _f(s, "4522", "completed_on") == "8/1/26", json.dumps(res)[:200])
    res = s.call("update_project", project_no="4522", company_id="acme",
                 fields={"completed_on": "2026-08-09"})
    r.check("today is allowed (the frozen today, not the wall clock)",
            res.get("ok") is True, json.dumps(res)[:200])
    res = s.call("update_project", project_no="4521", company_id="acme",
                 fields={"completed_on": None})
    r.check("null clears it (Reopen)",
            res.get("ok") is True and _f(s, "4521", "completed_on") is None
            and "completed_on" in (_rec(s, "4521") or {}), json.dumps(res)[:200])
    r.check("and reopening leaves the bucket where it was",
            _f(s, "4521", "tracker_status") == "action_admin")
    res = s.call("create_project", fields={"project_no": "4530", "company_id": "acme",
                                           "completed_on": "2026-08-02"})
    r.check("create_project carries completed_on",
            res.get("ok") is True and _f(s, "4530", "completed_on") == "2026-08-02",
            json.dumps(res)[:200])

    # ---- 2. refused ------------------------------------------------------------
    r.section("unreadable and future dates are refused, naming the values")
    seed()
    before = json.dumps(_rec(s, "4524"), sort_keys=True)
    for bad in ("soon", "2026-13-45", "2026-9-1junk", "2026-9- 1", ""):
        res = s.call("update_project", project_no="4524", company_id="acme",
                     fields={"completed_on": bad})
        r.check(f"completed_on {bad!r} is refused, naming the field and the value",
                res.get("ok") is False and "completed_on" in str(res.get("error"))
                and repr(bad) in str(res.get("error")), json.dumps(res)[:200])
    for bad in (20260801, True, ["2026-08-01"]):
        res = s.call("update_project", project_no="4524", company_id="acme",
                     fields={"completed_on": bad})
        r.check(f"completed_on {json.dumps(bad)} (wrong type) is refused",
                res.get("ok") is False and "completed_on" in str(res.get("error")),
                json.dumps(res)[:200])
    res = s.call("update_project", project_no="4524", company_id="acme",
                 fields={"completed_on": "2026-08-10"})
    err = str(res.get("error"))
    r.check("tomorrow is refused, naming both dates",
            res.get("ok") is False and "2026-08-10" in err and "2026-08-09" in err,
            json.dumps(res)[:200])
    res = s.call("update_project", project_no="4524", company_id="acme",
                 fields={"completed_on": "8/20/26"})
    r.check("a future date is refused in tracker style too, though the wall clock "
            "has passed it", res.get("ok") is False and "8/20/26" in str(res.get("error")),
            json.dumps(res)[:200])
    res = s.call("create_project", fields={"project_no": "4531", "company_id": "acme",
                                           "completed_on": "2026-08-10"})
    r.check("create_project refuses a future completion as well",
            res.get("ok") is False and _rec(s, "4531") is None, json.dumps(res)[:200])
    r.check("the record is byte-identical after every refusal",
            json.dumps(_rec(s, "4524"), sort_keys=True) == before)

    # ---- 3. before the deal date: allowed, with a warning -----------------------
    r.section("before the deal date is allowed with a warning")
    seed()
    res = s.call("update_project", project_no="4521", company_id="acme",
                 fields={"completed_on": "2026-03-01"})
    warn = " ".join(str(w) for w in res.get("warnings") or [])
    r.check("a completion before the deal date is saved",
            res.get("ok") is True and _f(s, "4521", "completed_on") == "2026-03-01",
            json.dumps(res)[:200])
    r.check("and the response warns, naming both dates",
            "2026-03-01" in warn and "2026-03-14" in warn, json.dumps(res.get("warnings")))
    res = s.call("update_project", project_no="4522", company_id="acme",
                 fields={"completed_on": "2026-03-01"})
    warn = " ".join(str(w) for w in res.get("warnings") or [])
    r.check("a tracker-style deal date is compared as a date, not as text",
            res.get("ok") is True and "3/14/2026" in warn, json.dumps(res)[:200])
    res = s.call("update_project", project_no="4523", company_id="acme",
                 fields={"completed_on": "2026-03-01"})
    r.check("an unreadable deal date never blocks the completion, and draws no warning",
            res.get("ok") is True and not res.get("warnings"), json.dumps(res)[:200])
    res = s.call("update_project", project_no="4524", company_id="acme",
                 fields={"completed_on": "2026-03-01"})
    r.check("no deal date at all: saved, no warning",
            res.get("ok") is True and not res.get("warnings"), json.dumps(res)[:200])
    res = s.call("update_project", project_no="4521", company_id="acme",
                 fields={"notes": "unrelated edit"})
    r.check("an unrelated edit does not repeat the warning", res.get("ok") is True
            and not res.get("warnings"), json.dumps(res)[:200])
    res = s.call("create_project", fields={"project_no": "4532", "company_id": "acme",
                                           "date": "2026-05-01", "completed_on": "2026-04-01"})
    warn = " ".join(str(w) for w in res.get("warnings") or [])
    r.check("create_project warns the same way",
            res.get("ok") is True and "2026-04-01" in warn and "2026-05-01" in warn,
            json.dumps(res)[:200])

    # ---- 3b. next_action_due: a completed job is not due (G2) --------------------
    r.section("next_action_due drops a completed project, the Live screen's rule")
    s.reset(companies=[company("acme", "Ace Manufacturing"), company("beta", "Beta Works")],
            projects=[project("4521", "acme", next_action_on="2026-08-01"),
                      project("4521", "beta", next_action_on="2026-08-01"),
                      project("4540", "acme", next_action_on="2026-08-01", completed_on=""),
                      project("4541", "acme", next_action_on="2026-08-01", completed_on="  "),
                      project("4542", "acme", next_action_on="2026-08-01", completed_on="8/2/26")])

    def due():
        res = s.call("list_projects", next_action_due=True)
        return sorted((str(p["project_no"]), p["company_id"]) for p in res.get("projects", []))
    r.check("before: every past-due next action is due",
            due() == [("4521", "acme"), ("4521", "beta"), ("4540", "acme"), ("4541", "acme")],
            json.dumps(due()))
    s.call("update_project", project_no="4521", company_id="acme", fields={"completed_on": "2026-08-05"})
    r.check("a project marked complete drops out of next_action_due",
            ("4521", "acme") not in due(), json.dumps(due()))
    r.check("... and the other customer's 4521 is still due", ("4521", "beta") in due(), json.dumps(due()))
    r.check("a completed_on of '' or blanks is not a completion (the view reads it the same way)",
            ("4540", "acme") in due() and ("4541", "acme") in due(), json.dumps(due()))
    r.check("a tracker-style completed_on on disk is a completion", ("4542", "acme") not in due(),
            json.dumps(due()))
    s.call("update_project", project_no="4521", company_id="acme", fields={"completed_on": None})
    r.check("reopened, it is due again", ("4521", "acme") in due(), json.dumps(due()))
    r.check("without the flag, completed projects are still listed",
            len(s.call("list_projects").get("projects", [])) == 5)

    # ---- 4a. a fresh process reads it back -----------------------------------------
    r.section("round trip through a fresh process")
    seed()
    s.call("update_project", project_no="4521", company_id="acme",
           fields={"completed_on": "2026-08-01"})
    probe = (
        "import sys, json, asyncio\nfrom pathlib import Path\n"
        f"sys.path.insert(0, {str(Path(__file__).resolve().parents[1])!r})\n"
        "from lib.harness import load_server\n"
        f"srv = load_server({str(crm)!r})\n"
        f"srv.STORE = srv.Store(Path({str(s.path)!r}))\n"
        "res = asyncio.run(srv.mcp.call_tool('get_project', "
        "{'project_no': '4521', 'company_id': 'acme'}))\n"
        "out = res[1]['result'] if isinstance(res, tuple) else json.loads(res[0].text)\n"
        "print(json.dumps(out.get('project', {}).get('completed_on')))\n")
    pr = subprocess.run([sys.executable, "-c", probe], capture_output=True, text=True, timeout=60)
    r.check("a separate process reads completed_on off disk",
            pr.stdout.strip().splitlines()[-1:] == ['"2026-08-01"'],
            (pr.stdout + pr.stderr)[-200:])
    log = [json.loads(l) for l in (s.path / "changelog.jsonl").read_text().splitlines()
           if l.strip()] if (s.path / "changelog.jsonl").exists() else []
    r.check("the changelog records it under the project, with its customer",
            any(e.get("entity") == "project" and e.get("key") == "4521"
                and e.get("company_id") == "acme"
                and (e.get("fields") or {}).get("completed_on") == "2026-08-01" for e in log),
            json.dumps(log)[-200:])

    # ---- 4b. a re-import through the operator's path ----------------------------------
    r.section("a re-import keeps completed_on while the importer re-sets the bucket")
    _reimport_checks(r, srv, crm)

    # ---- 5. two customers, one number ---------------------------------------------------
    r.section("two customers holding 4521")
    s.reset(companies=[company("acme", "Ace Manufacturing"), company("beta", "Beta Works")],
            projects=[project("4521", "acme", tracker_status="action_admin"),
                      project("4521", "beta", tracker_status="action_owner")])
    beta_before = json.dumps(_rec(s, "4521", "beta"), sort_keys=True)
    res = s.call("update_project", project_no="4521", company_id="acme",
                 fields={"completed_on": "2026-08-01"})
    r.check("completing Ace Manufacturing's 4521 succeeds",
            res.get("ok") is True and _f(s, "4521", "completed_on") == "2026-08-01",
            json.dumps(res)[:200])
    r.check("and Beta Works' 4521 is untouched",
            json.dumps(_rec(s, "4521", "beta"), sort_keys=True) == beta_before)
    merge = _load(crm, "merge.py", "_cmp_merge")
    both_before = json.dumps(sorted((json.dumps(p, sort_keys=True) for p in s.read("projects"))))
    merged, rep = merge.merge_all(
        {"companies.json": [company("acme", "Ace Manufacturing"), company("beta", "Beta Works")],
         "contacts.json": [], "vendors.json": [], "shipments.json": [], "invoices.json": [],
         "needs_review.json": [],
         "projects.json": [project("4521", "acme", revenue=99, tracker_status="action_owner")]},
        str(s.path))
    after = json.dumps(sorted(json.dumps(p, sort_keys=True) for p in merged["projects.json"]))
    r.check("a re-import leaves both 4521s exactly as they were (the ambiguous path)",
            after == both_before, after[:200])
    r.check("and reports the number as ambiguous",
            any(a.get("key") == "4521" for a in rep.get("ambiguous", [])), json.dumps(rep)[:200])
    return r


def _load(crm, name, as_name):
    import importlib.util
    spec = importlib.util.spec_from_file_location(as_name, Path(crm) / "pipeline" / name)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


def _reimport_checks(r, srv, crm):
    """normalize.run(mode="merge") -- the path the operator's import takes -- on
    the Live Tracker fixture workbook. Project 5001 is held by one customer
    only, so the ordinary refresh path runs, not the ambiguous one."""
    import openpyxl
    from openpyxl.styles import PatternFill
    import test_livetracker as lt
    tmp = Path(tempfile.mkdtemp(prefix="crm-completed-"))
    added = str(Path(crm) / "pipeline")
    sys.path.insert(0, added)       # normalize imports merge as a sibling
    try:
        nrm = lt._load(crm, "normalize.py")
        xl = tmp / "wb.xlsx"
        lt._build_workbook(xl)
        store = tmp / "store"
        store.mkdir()
        nrm.run(str(xl), str(store), force=False, mode="merge")

        def rec():
            ps = [p for p in json.loads((store / "projects.json").read_text())
                  if str(p.get("project_no")) == "5001"]
            return ps[0] if len(ps) == 1 else None
        p = rec()
        if not r.check("the fixture holds 5001 once, bucketed by its colour",
                       p is not None and p.get("tracker_status") == "action_admin",
                       json.dumps(p)[:200]):
            return
        s = Store.__new__(Store)
        s.server, s.path = srv, store
        s.rebind()
        res = s.call("update_project", project_no="5001", company_id=p["company_id"],
                     fields={"completed_on": "2026-08-01"})
        r.check("5001 is marked complete", res.get("ok") is True, json.dumps(res)[:200])
        # the sheet recolours the row: the importer owns the bucket
        wb = openpyxl.load_workbook(xl)
        wb["Project Tracker"].cell(row=2, column=6).fill = PatternFill(
            start_color=YELLOW, end_color=YELLOW, fill_type="solid")
        wb.save(xl)
        nrm.run(str(xl), str(store), force=False, mode="merge")
        p = rec() or {}
        r.check("after a re-import completed_on is still set",
                p.get("completed_on") == "2026-08-01", json.dumps(p)[:200])
        r.check("while the importer re-set tracker_status from the new colour",
                p.get("tracker_status") == "action_owner", json.dumps(p)[:200])
        # Reopen, re-import: the null the operator wrote holds too
        s.rebind()
        s.call("update_project", project_no="5001", company_id=p.get("company_id"),
               fields={"completed_on": None})
        nrm.run(str(xl), str(store), force=False, mode="merge")
        p = rec() or {}
        r.check("reopened, then re-imported: still reopened",
                "completed_on" in p and p.get("completed_on") is None, json.dumps(p)[:200])
    finally:
        try:
            sys.path.remove(added)
        except ValueError:
            pass
        shutil.rmtree(tmp, ignore_errors=True)
