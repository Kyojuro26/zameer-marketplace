"""The quote pipeline (0.1.41, Phase E).

A quote request that sits unanswered is invisible today -- the operator lost
a ~$50K deal to two extra options that took ~1.5 weeks. Decisions asserted,
each with an obvious wrong version:

 1. FIELDS: quote_requested_on, quote_sent_on and quote_revisions
    ([{requested_on, sent_on, note}]) on a project. A sent date before its
    requested date is refused naming both -- checked against the record as it
    will be saved, not just the fields sent. A revision needs requested_on.
    Dates pass the same grammar next_action_on does. Nothing is backfilled:
    an import never writes them, and a re-import keeps what the operator set.
 2. THE REPORT: crm_metrics(report="quotes") -- three ranked lists, oldest
    first, ages in business days (Mon-Fri):
      waiting_to_send         a request (or an open revision) with no sent date;
                              flagged past the SLA
      sent_awaiting_decision  pending projects with a sent date; flagged when no
                              next_action_on lies in the future
      stale_pending           pending with no activity for more than N days --
                              listed, never changed
    quote_turnaround_days is the median business days request -> sent over the
    first send and every revision, with its denominator; win_rate is
    won / (won + lost) over decided projects only, and says pending are
    undecided and many may be dead.
 3. SETTINGS: the SLA (2 business days) and N (60 days) live in the store's
    settings.json, changed with update_store_settings -- validated, logged,
    never in code.

Names are invented.
"""
import datetime
import json
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from lib.harness import Result, Store, company, load_server, project  # noqa: E402

TODAY = datetime.date(2026, 4, 15)          # a Wednesday


def P(no, status="pending", **kw):
    return project(no, "acme", status=status, **kw)


def seed(st):
    st.reset(
        companies=[company("acme", "Ace Manufacturing")],
        projects=[
            P("5001", date="2026-04-01", quote_requested_on="2026-04-09"),   # Thu: 4 bd
            P("5002", date="2026-04-01", quote_requested_on="2026-04-13"),   # Mon: 2 bd
            P("5003", date="2026-03-01", quote_requested_on="2026-03-02",
              quote_sent_on="2026-03-04",                                    # 2 bd
              quote_revisions=[
                  {"requested_on": "2026-03-20", "sent_on": "2026-03-31",    # 7 bd
                   "note": "two more options"},
                  {"requested_on": "2026-04-10", "note": "a third option"}]),  # open, 3 bd
            P("5004", date="2026-03-01", quote_requested_on="2026-03-02",
              quote_sent_on="2026-03-03", next_action_on="2026-04-20",       # 1 bd
              next_action="call back"),
            P("5005", date="2026-01-20", quote_sent_on="2026-02-02"),        # no request date
            P("5006", date="2026-01-05"),                                    # 100 days, nothing
            P("5007", date="2026-04-01"),
            P("5012"),                                                       # no date at all
            P("5008", status="won", date="2025-12-01", revenue=5000, total_cost=3000,
              quote_requested_on="2026-01-05", quote_sent_on="2026-01-09"),  # 4 bd
            P("5009", status="lost", date="2025-11-01"),
            P("5010", status="lost", date="2025-11-02"),
            P("5011", date="2025-12-01"),                                    # old, but edited
        ])


def quotes(st):
    got = st.call("crm_metrics", report="quotes")
    return (got.get("reports") or {}).get("quotes") or {}, got


def run(server, crm_dir=None):
    r = Result("quotes", since="0.1.41")
    crm = Path(crm_dir) if crm_dir else \
        Path(__file__).resolve().parents[2] / "plugins/unrivaled-solutions/skills/crm"
    if server is None:
        server = load_server(crm)
    hooked = callable(getattr(server, "_today", None))
    real_today = getattr(server, "_today", None)
    if hooked:
        server._today = lambda: TODAY
    tmp = Path(tempfile.mkdtemp(prefix="crmquote-"))
    try:
        _body(r, server, crm, tmp)
    finally:
        if hooked:
            server._today = real_today
        shutil.rmtree(tmp, ignore_errors=True)
    return r


def _nos(rows):
    return [(x.get("project_no"), x.get("kind")) for x in rows or []]


def _body(r, server, crm, tmp):
    st = Store(server, tmp / "m" / "store")
    seed(st)
    # an edit is activity: 5011 is old by its dates, but the operator touched it
    st.call("update_project", project_no="5011", fields={"notes": "chased by phone"})

    # ---- 2. the report ------------------------------------------------------
    r.section("the quotes report")
    q, got = quotes(st)
    if not r.check("crm_metrics(report='quotes') answers", got.get("ok") is True and bool(q), got):
        return
    w = q.get("waiting_to_send") or {}
    r.check("waiting to send, oldest first: 5001, 5003's open revision, 5002",
            _nos(w.get("rows")) == [("5001", "quote"), ("5003", "revision"), ("5002", "quote")],
            _nos(w.get("rows")))
    ages = {(x.get("project_no"), x.get("kind")): x for x in w.get("rows") or []}
    r.check("ages in business days: 4, 3 (a Friday request), 2",
            [ages.get(k, {}).get("age_business_days") for k in
             (("5001", "quote"), ("5003", "revision"), ("5002", "quote"))] == [4, 3, 2],
            [x.get("age_business_days") for x in w.get("rows") or []])
    r.check("past the 2-business-day SLA: 5001 and the revision, not 5002",
            [ages.get(k, {}).get("past_sla") for k in
             (("5001", "quote"), ("5003", "revision"), ("5002", "quote"))] == [True, True, False],
            [x.get("past_sla") for x in w.get("rows") or []])
    r.check("the revision row carries its note", ages.get(("5003", "revision"), {}).get("note")
            == "a third option", ages.get(("5003", "revision")))
    s = q.get("sent_awaiting_decision") or {}
    r.check("sent, awaiting decision, oldest send first: 5005, 5004, 5003 (its last send)",
            [x.get("project_no") for x in s.get("rows") or []] == ["5005", "5004", "5003"],
            [x.get("project_no") for x in s.get("rows") or []])
    sa = {x.get("project_no"): x for x in s.get("rows") or []}
    r.check("... flagged when no follow-up lies in the future: 5005 and 5003, not 5004",
            [sa.get(n, {}).get("no_follow_up") for n in ("5005", "5004", "5003")]
            == [True, False, True], sa)
    r.check("... a won project is decided, not awaiting", "5008" not in sa)
    sp = q.get("stale_pending") or {}
    r.check("stale pending, oldest first: 5012 (no dated activity), 5006 (100 days), "
            "5005 (72 days); an edited project is not stale",
            [x.get("project_no") for x in sp.get("rows") or []] == ["5012", "5006", "5005"],
            [(x.get("project_no"), x.get("days_since_activity")) for x in sp.get("rows") or []])
    r.check("... ages are calendar days since the last activity",
            [x.get("days_since_activity") for x in sp.get("rows") or []] == [None, 100, 72],
            [x.get("days_since_activity") for x in sp.get("rows") or []])
    r.check("... and the report changed nothing: every one is still pending",
            all(p.get("status") == "pending" for p in st.read("projects")
                if p["project_no"] in ("5012", "5006", "5005")))
    t = q.get("quote_turnaround_days") or {}
    r.check("turnaround: median of 1, 2, 4, 7 business days = 3, measured on 4 of 5 quotes",
            t.get("value") == 3 and t.get("counted") == 4
            and t.get("excluded") == {"no_request_date": 1}, t)
    r.check("... and its basis says so", "4 of 5 quotes" in str(t.get("basis"))
            and "1 has no request date" in str(t.get("basis")), t.get("basis"))
    wr = q.get("win_rate") or {}
    r.check("win rate: 1 won / (1 won + 2 lost) over the 3 decided",
            wr.get("counted") == 3 and abs((wr.get("value") or 0) - 1 / 3) < 1e-9
            and wr.get("excluded") == {"not_decided": 9}, wr)
    r.check("... its basis is honest about the undecided",
            "won / (won + lost) over decided projects only" in str(wr.get("basis"))
            and "9 pending are undecided and many may be dead" in str(wr.get("basis")),
            wr.get("basis"))
    r.check("the report states the settings it used: SLA 2 business days, N 60 days",
            (q.get("settings") or {}) == {"quote_sla_business_days": 2, "stale_pending_days": 60},
            q.get("settings"))

    # ---- 3. settings ---------------------------------------------------------
    r.section("settings live in the store")
    got = st.call("update_store_settings",
                  fields={"quote_sla_business_days": 5, "stale_pending_days": 90})
    r.check("update_store_settings changes them", got.get("ok") is True, got)
    stored = json.loads((st.path / "settings.json").read_text()) \
        if (st.path / "settings.json").exists() else None
    r.check("... into the store's settings.json", stored == {"quote_sla_business_days": 5,
                                                            "stale_pending_days": 90}, stored)
    q2, _ = quotes(st)
    r.check("a 5-day SLA: nothing is past it now",
            not any(x.get("past_sla") for x in (q2.get("waiting_to_send") or {}).get("rows") or []))
    r.check("N = 90: 5005 (72 days) is no longer stale",
            [x.get("project_no") for x in (q2.get("stale_pending") or {}).get("rows") or []]
            == ["5012", "5006"], (q2.get("stale_pending") or {}).get("rows"))
    for label, f in (("zero days", {"quote_sla_business_days": 0}),
                     ("text", {"stale_pending_days": "sixty"}),
                     ("a boolean", {"quote_sla_business_days": True}),
                     ("an unknown setting", {"colour": "blue"})):
        before = (st.path / "settings.json").read_text()
        got = st.call("update_store_settings", fields=f)
        r.check(f"settings refuse {label}", got.get("ok") is False and "_raised" not in got, got)
        r.check(f"... and settings.json is unchanged ({label})",
                (st.path / "settings.json").read_text() == before)
    log = [json.loads(l) for l in (st.path / "changelog.jsonl").read_text().splitlines() if l.strip()]
    r.check("a settings change is logged", any(e.get("entity") == "settings" for e in log))
    st.call("update_store_settings", fields={"quote_sla_business_days": 2, "stale_pending_days": 60})

    # ---- 1. fields and validation --------------------------------------------
    r.section("the fields, validated")

    def refused(label, fields, needles=()):
        before = st.raw("projects")
        got = st.call("update_project", project_no="5001", fields=fields)
        r.check(f"refused: {label}", got.get("ok") is False and "_raised" not in got, got)
        for n in needles:
            r.check(f"... the refusal names {n}", n in str(got.get("error")), got.get("error"))
        r.check(f"... nothing written ({label})", st.raw("projects") == before)

    refused("a sent date before the stored requested date",
            {"quote_sent_on": "2026-04-01"}, ("2026-04-01", "2026-04-09"))
    refused("a requested date after the... send in the same edit",
            {"quote_requested_on": "2026-04-14", "quote_sent_on": "2026-04-13"},
            ("2026-04-13", "2026-04-14"))
    refused("a date the screen cannot read", {"quote_requested_on": "next week"})
    refused("revisions that are not a list", {"quote_revisions": {"requested_on": "2026-04-01"}})
    refused("a revision with no requested_on", {"quote_revisions": [{"note": "more options"}]})
    refused("a revision with an unknown field",
            {"quote_revisions": [{"requested_on": "2026-04-10", "who": "Rae"}]})
    refused("a revision sent before it was requested",
            {"quote_revisions": [{"requested_on": "2026-04-10", "sent_on": "2026-04-09"}]},
            ("2026-04-09", "2026-04-10"))
    got = st.call("update_project", project_no="5001", fields={"quote_sent_on": "2026-04-15"})
    r.check("marking 5001 sent today is accepted", got.get("ok") is True, got)
    q3, _ = quotes(st)
    r.check("... it leaves waiting-to-send and joins sent-awaiting",
            ("5001", "quote") not in _nos((q3.get("waiting_to_send") or {}).get("rows"))
            and "5001" in [x.get("project_no") for x in
                           (q3.get("sent_awaiting_decision") or {}).get("rows") or []])
    r.check("... and the turnaround now measures 5 of 6 (it took 4 business days)",
            (q3.get("quote_turnaround_days") or {}).get("counted") == 5,
            q3.get("quote_turnaround_days"))

    # ---- adversarial: nothing waiting may drop out of sight ------------------
    r.section("an unreadable request date, and a number two customers share")
    adv = Store(server, tmp / "adv" / "store")
    adv.reset(companies=[company("acme", "Ace Manufacturing"), company("beta", "Beta Works")],
              projects=[P("6001", date="2026-04-01", quote_requested_on="someday"),
                        project("6003", "acme", status="pending", date="2026-01-01"),
                        project("6003", "beta", status="pending", date="2026-01-01")])
    adv.call("update_project", project_no="6003", company_id="beta", fields={"notes": "x"})
    qa, _ = quotes(adv)
    wa = {(x.get("project_no"), x.get("kind")): x for x in
          (qa.get("waiting_to_send") or {}).get("rows") or []}
    r.check("a request date nobody can read is still listed as waiting, never dropped",
            ("6001", "quote") in wa and wa[("6001", "quote")].get("date_unreadable") is True
            and wa[("6001", "quote")].get("age_business_days") is None, wa)
    stl = [(x.get("project_no"), x.get("company_id")) for x in
           (qa.get("stale_pending") or {}).get("rows") or []]
    r.check("an edit to Beta's 6003 is not activity on Acme's 6003",
            ("6003", "acme") in stl and ("6003", "beta") not in stl, stl)

    # ---- a broken settings file costs the Quotes page only --------------------
    r.section("a bad settings.json")
    bad = Store(server, tmp / "bad" / "store")
    seed(bad)
    (bad.path / "settings.json").write_text('{"quote_sla_business_days": "two"}')
    got = bad.call("crm_metrics", report="quotes")
    r.check("the quotes report refuses in a sentence naming the setting",
            got.get("ok") is False and "quote_sla_business_days" in str(got.get("error")), got)
    import subprocess
    out = subprocess.run([sys.executable, str(Path(crm) / "view" / "build_view.py"),
                          "--store", str(bad.path), "--out", str(tmp / "bad.html")],
                         capture_output=True, text=True)
    html = (tmp / "bad.html").read_text() if (tmp / "bad.html").exists() else ""
    said = out.stdout + out.stderr
    r.check("... the page still builds, with the company figures and the CFO report",
            '"exposure_open_receivable_usd"' in html and '"cfo": {' in html, said[-300:])
    r.check("... and the warning names the Quotes page, not the company figures",
            "quotes" in said.lower() and "metrics not embedded" not in said, said[-300:])

    # ---- nothing backfilled; a re-import keeps the operator's quote dates -----
    r.section("imports never write the quote fields")
    import importlib.util
    p = Path(crm) / "pipeline" / "normalize.py"
    spec = importlib.util.spec_from_file_location("_nrm_quotes", p)
    nrm = importlib.util.module_from_spec(spec)
    sys.path.insert(0, str(p.parent))
    try:
        spec.loader.exec_module(nrm)
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sales Tracker 2026"
        ws.append(["Project#", "Date", "Customer", "Description", "Location", "Status",
                   "PO Y/N", "Client PO#", "Invoice #", "Revenue", "Total Cost",
                   "Total GP", "Margain", "Notes"])
        ws.append(["9001", "2026-03-01", "Brightwater Fabrication", "Line install",
                   "Dayton OH", "pending", "N", None, None, 50000, 30000, 20000, 0.4, ""])
        pt = wb.create_sheet("Project Tracker")
        pt.append(["Unrivaled Project#:", "Client PO#:", "Start Date:", "Client Name:",
                   "Client Location:", "Open Orders Notes:", "Vendor 1 PO#:",
                   "Vendor 1 Ship Date:"])
        wb.create_sheet("Client Contacts").append(["Client Business", "Client Name", "Email"])
        wb.create_sheet("Vendor Contacts").append(["Company", "Headquarters Location"])
        x = tmp / "tracker.xlsx"
        wb.save(x)
        imp = Store(server, tmp / "imp" / "store")
        nrm.run(str(x), str(imp.path), force=True, mode="merge")
        imp.rebind()
        rec = next((p_ for p_ in imp.read("projects") if p_.get("project_no") == "9001"), {})
        r.check("an imported pending project carries no quote dates -- its project date "
                "is not a request date", rec.get("date") and not rec.get("quote_requested_on")
                and not rec.get("quote_sent_on") and not rec.get("quote_revisions"), rec)
        imp.call("update_project", project_no="9001",
                 fields={"quote_requested_on": "2026-03-02",
                         "quote_revisions": [{"requested_on": "2026-03-05", "note": "x"}]})
        nrm.run(str(x), str(imp.path), force=True, mode="merge")
        rec = next((p_ for p_ in imp.read("projects") if p_.get("project_no") == "9001"), {})
        r.check("a re-import keeps the quote dates the operator set",
                rec.get("quote_requested_on") == "2026-03-02"
                and rec.get("quote_revisions") == [{"requested_on": "2026-03-05", "note": "x"}], rec)
    except Exception as e:                            # noqa: BLE001
        r.check("the import fixture ran", False, f"{type(e).__name__}: {e}")
    finally:
        sys.path.remove(str(p.parent))
