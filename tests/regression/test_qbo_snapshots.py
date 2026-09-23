"""QuickBooks export parsers and the snapshot layer they feed (0.1.38, B1).

Decisions asserted, each with an obvious wrong version:

 1. COLUMNS ARE FOUND BY HEADER NAME. A file whose columns are reordered
    parses to the same rows; a file whose header row does not match is
    refused WHOLE, the refusal carries the header that was found, and no rows
    come back. Reading by position files the open balance as the amount the
    first time QuickBooks moves a column.
 2. MONEY IS INTEGER CENTS, EXACTLY. 0.29 is 29, never 28; an amount that is
    not a whole number of cents is refused, not rounded.
 3. THE WINDOW COMES FROM TITLE ROW 3, as_of FROM THE FOOTER. "January
    1-March 31, 2026" is 2026-01-01..2026-03-31, and a window QuickBooks spans
    across a year keeps both years.
 4. THE VENDOR LIST IS GROUPED. A transaction belongs to the vendor row above
    it; "Total for ..." and the grand TOTAL are never transactions; a blank
    split account is null, never guessed from the account column.
 5. A SNAPSHOT IS A SIBLING OF THE STORE, never inside it, replaced wholesale,
    written through the store's own Windows-lock retry, and carries kind,
    source, as_of, window and loaded_at. Neither the store's startup nor the
    importer reads or writes it.
 6. THE PARSERS ARE READ-ONLY on their input.

Fixtures copy the export LAYOUT only. Every name and figure is invented: the
real exports carry customer and vendor names and must never enter this repo.
"""
import json
import os
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from lib.harness import Result, Store, load_server  # noqa: E402

INV_HEADER = ["Date", "Transaction type", "Num", "Name", "Memo", "Due date",
              "Amount", "Open balance"]
TXN_HEADER = [None, "Date", "Track 1099", "Transaction type", "Num",
              "Posting (Y/N)", "Memo", "Account full name",
              "Item split account", "Amount"]
FOOTER = " Tuesday, April 7, 2026 09:15 AM GMT-04:00"


def _load(crm, name):
    import importlib.util
    p = Path(crm) / "pipeline" / f"{name}.py"
    if not p.exists():
        return None
    spec = importlib.util.spec_from_file_location(f"_{name}_under_test", p)
    m = importlib.util.module_from_spec(spec)
    sys.path.insert(0, str(p.parent))
    try:
        spec.loader.exec_module(m)
    finally:
        sys.path.remove(str(p.parent))
    return m


def _xlsx(path, rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for row in rows:
        ws.append(list(row))
    wb.save(path)
    return path


def _titled(title, window, header, body, footer=FOOTER):
    width = len(header)
    pad = [None] * (width - 1)
    rows = [["Invented Books LLC"] + pad, [title] + pad, [window] + pad,
            [None] * width, header]
    rows += body
    rows += [[None] * width, [None] * width]
    if footer is not None:
        rows.append([footer] + pad)
    return rows


def _reorder(header, rows, perm):
    """The same table with its columns permuted: header and every row."""
    return [header[i] for i in perm], [[r[i] for i in perm] for r in rows]


INV_BODY = [
    ["01/05/2026", "Invoice", "7001", "Brightwater Fabrication", None,
     "02/04/2026", 12500.0, 12500],
    ["02/10/2026", "Invoice", "7002", "Ironvale Supply", "Line install",
     "03/12/2026", 0.29, 0.0],
    ["03/31/2026", "Invoice", "7003", "Brightwater Fabrication", None,
     "04/30/2026", 1571990.42, 135327.08],
]
INV_TOTAL = ["TOTAL", None, None, None, None, None, 1584490.71, 147827.08]

TXN_BODY = [
    ["Cobalt Freight", None, None, None, None, None, None, None, None, None],
    [None, "01/08/2026", "No", "Purchase Order", "1167", "No", None,
     "Accounts Payable", None, 4200.0],
    [None, "01/20/2026", "No", "Bill", "INV-88", "Yes", None,
     "Accounts Payable", "Cost of Goods Sold", 4200.0],
    [None, "02/02/2026", "No", "Bill Payment (Check)", "3301", "Yes", None,
     None, "Checking", -4200],
    ["Total for Cobalt Freight", None, None, None, None, None, None, None,
     None, 4200.0],
    ["Trueline Metals", None, None, None, None, None, None, None, None, None],
    [None, "03/03/2026", "Yes", "Expense", "", "Yes", "plate stock",
     "Checking", "Cost of Goods Sold", 815.55],
    [None, "03/04/2026", "No", "Journal Entry", "JE-7", "Yes", None, None,
     None, None],
    ["Total for Trueline Metals", None, None, None, None, None, None, None,
     None, 815.55],
    ["TOTAL", None, None, None, None, None, None, None, None, 5015.55],
]


def run(server, crm_dir=None):
    r = Result("qbo-snapshots", since="0.1.38")
    crm = Path(crm_dir) if crm_dir else \
        Path(__file__).resolve().parents[2] / "plugins/unrivaled-solutions/skills/crm"
    if server is None:                    # the mutation runner passes None
        server = load_server(crm)
    tmp = Path(tempfile.mkdtemp(prefix="qbotest-"))
    qx = _load(crm, "qbo_exports")
    r.check("pipeline/qbo_exports.py exists", qx is not None,
            "no parser -- an export has no way into the snapshot")

    try:
        if qx is not None:
            _parsers(r, qx, tmp)
        _snapshots(r, server, crm, tmp)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return r


# ---------------------------------------------------------------- parsers --

def _refusal(qx, fn, path):
    """(error, rows) -- rows is whatever came back if nothing was raised."""
    try:
        out = fn(str(path))
    except qx.QboExportError as e:
        return e, None
    except Exception as e:                            # noqa: BLE001
        # a crash is not a refusal: report it as "not refused", so the named
        # check that expected a refusal is the one that fails
        return None, {"_crashed": f"{type(e).__name__}: {e}"}
    return None, out


def _parsers(r, qx, tmp):
    # ---- 1-3. the invoice list -----------------------------------------------
    r.section("invoice list: header, cents, window, as_of")
    inv = _xlsx(tmp / "inv.xlsx", _titled(
        "Invoice List by Date", "January 1-March 31, 2026", INV_HEADER,
        INV_BODY + [INV_TOTAL]))
    before = (inv.read_bytes(), os.stat(inv).st_mtime_ns)
    err, got = _refusal(qx, qx.parse_invoice_list, inv)
    if not r.check("a well-formed invoice list parses", err is None, str(err)):
        return
    r.check("the parser leaves its input byte-identical and unmodified",
            (inv.read_bytes(), os.stat(inv).st_mtime_ns) == before)
    r.check("kind is invoices", got.get("kind") == "invoices", got.get("kind"))
    r.check("window from title row 3",
            (got.get("window_start"), got.get("window_end"))
            == ("2026-01-01", "2026-03-31"),
            f"{got.get('window_start')}..{got.get('window_end')}")
    r.check("as_of is the footer's timestamp",
            got.get("as_of") == "2026-04-07T09:15:00-04:00", got.get("as_of"))
    rows = got.get("rows") or []
    r.check("three invoices: TOTAL, blanks and the footer are not invoices",
            len(rows) == 3, f"{len(rows)} rows")
    by = {x.get("num"): x for x in rows}
    r.check("an invoice row carries type and Num as the key pair",
            by.get("7001", {}).get("type") == "Invoice", by.get("7001"))
    r.check("amounts are integer cents: 12500.0 -> 1250000",
            by.get("7001", {}).get("amount_cents") == 1250000
            and type(by["7001"]["amount_cents"]) is int, by.get("7001"))
    r.check("0.29 is 29 cents, not 28 (float truncation)",
            by.get("7002", {}).get("amount_cents") == 29, by.get("7002"))
    r.check("1571990.42 is 157199042 cents exactly",
            by.get("7003", {}).get("amount_cents") == 157199042, by.get("7003"))
    r.check("open balance 135327.08 -> 13532708 cents, a separate field",
            by.get("7003", {}).get("open_cents") == 13532708, by.get("7003"))
    r.check("an int open balance (12500) is cents too",
            by.get("7001", {}).get("open_cents") == 1250000
            and type(by["7001"]["open_cents"]) is int, by.get("7001"))
    r.check("dates are ISO", (by.get("7001", {}).get("date"),
                              by.get("7001", {}).get("due_date"))
            == ("2026-01-05", "2026-02-04"), by.get("7001"))
    r.check("the customer name is kept as text, not a key",
            by.get("7002", {}).get("name") == "Ironvale Supply", by.get("7002"))

    r.section("invoice list: parsed by name, not position")
    perm = [7, 6, 5, 4, 3, 2, 1, 0]
    h2, b2 = _reorder(INV_HEADER, INV_BODY, perm)
    inv2 = _xlsx(tmp / "inv2.xlsx", _titled(
        "Invoice List by Date", "January 1-March 31, 2026", h2, b2))
    err, got2 = _refusal(qx, qx.parse_invoice_list, inv2)
    r.check("columns reversed -> the same rows",
            err is None and got2.get("rows") == rows,
            str(err) if err else "rows differ: read by position")

    r.section("invoice list: a wrong header is refused whole")
    bad_h = INV_HEADER[:7] + ["Balance"]
    inv3 = _xlsx(tmp / "inv3.xlsx", _titled(
        "Invoice List by Date", "January 1-March 31, 2026", bad_h, INV_BODY))
    err, got3 = _refusal(qx, qx.parse_invoice_list, inv3)
    r.check("a renamed column is refused, not loaded partially",
            err is not None and got3 is None, f"returned {got3!r}")
    r.check("the refusal carries the header it found",
            err is not None and list(getattr(err, "found_header", None) or [])
            == bad_h, getattr(err, "found_header", None))
    err, _ = _refusal(qx, qx.parse_vendor_transactions, inv)
    r.check("an invoice list fed to the vendor parser is refused",
            err is not None, "parsed as vendor transactions")
    r.check("... and that refusal names the invoice header",
            err is not None and list(getattr(err, "found_header", None) or [])
            == INV_HEADER, getattr(err, "found_header", None))

    r.section("invoice list: money that is not whole cents is refused")
    inv4 = _xlsx(tmp / "inv4.xlsx", _titled(
        "Invoice List by Date", "January 1-March 31, 2026", INV_HEADER,
        [INV_BODY[0][:6] + [10.005, 0]]))
    err, got4 = _refusal(qx, qx.parse_invoice_list, inv4)
    r.check("10.005 is refused, never rounded to 1000 or 1001",
            err is not None and got4 is None, f"returned {got4!r}")

    r.section("invoice list: windows")
    for text, want in (("December 1, 2025-January 31, 2026",
                        ("2025-12-01", "2026-01-31")),
                       ("September 1-17, 2026", ("2026-09-01", "2026-09-17"))):
        p = _xlsx(tmp / "w.xlsx", _titled("Invoice List by Date", text,
                                          INV_HEADER, INV_BODY))
        err, g = _refusal(qx, qx.parse_invoice_list, p)
        r.check(f"window {text!r} -> {want[0]}..{want[1]}",
                err is None and (g.get("window_start"), g.get("window_end")) == want,
                str(err) if err else f"{g.get('window_start')}..{g.get('window_end')}")
    p = _xlsx(tmp / "w.xlsx", _titled("Invoice List by Date", "All Dates",
                                      INV_HEADER, INV_BODY))
    err, g = _refusal(qx, qx.parse_invoice_list, p)
    r.check("a window it cannot read is refused (no window, no 'outside')",
            err is not None, f"returned window {g and g.get('window_start')}")

    # ---- 4. the vendor transaction list --------------------------------------
    r.section("vendor transactions: groups, totals, split account")
    txn = _xlsx(tmp / "txn.xlsx", _titled(
        "Transaction List by Vendor", "January 1-March 31, 2026", TXN_HEADER,
        TXN_BODY))
    before = (txn.read_bytes(), os.stat(txn).st_mtime_ns)
    err, got = _refusal(qx, qx.parse_vendor_transactions, txn)
    if not r.check("a well-formed vendor list parses", err is None, str(err)):
        return
    r.check("the vendor parser leaves its input unmodified",
            (txn.read_bytes(), os.stat(txn).st_mtime_ns) == before)
    r.check("kind is vendor_transactions",
            got.get("kind") == "vendor_transactions", got.get("kind"))
    rows = got.get("rows") or []
    r.check("five transactions: vendor rows and every Total are skipped",
            len(rows) == 5, f"{len(rows)} rows: {[x.get('type') for x in rows]}")
    types = [x.get("type") for x in rows]
    r.check("no 'Total for' or TOTAL row became a transaction",
            not any("total" in str(x.get("vendor") or "").lower()
                    or x.get("type") is None for x in rows), types)
    po = next((x for x in rows if x.get("type") == "Purchase Order"), {})
    r.check("a transaction belongs to the vendor row above it",
            po.get("vendor") == "Cobalt Freight"
            and next((x for x in rows if x.get("type") == "Expense"), {})
            .get("vendor") == "Trueline Metals",
            [(x.get("vendor"), x.get("type")) for x in rows])
    r.check("a blank split account is null, not the account column",
            "split_account" in po and po["split_account"] is None, po)
    bill = next((x for x in rows if x.get("type") == "Bill"), {})
    r.check("a present split account is kept",
            bill.get("split_account") == "Cost of Goods Sold", bill)
    r.check("Num keeps its text (a bill Num is the vendor's own number)",
            bill.get("num") == "INV-88" and po.get("num") == "1167",
            (bill.get("num"), po.get("num")))
    r.check("Posting (Y/N) is a boolean",
            po.get("posting") is False and bill.get("posting") is True,
            (po.get("posting"), bill.get("posting")))
    exp = next((x for x in rows if x.get("type") == "Expense"), {})
    r.check("815.55 -> 81555 cents", exp.get("amount_cents") == 81555, exp)
    pay = next((x for x in rows if x.get("type") == "Bill Payment (Check)"), {})
    r.check("a negative int amount -> -420000 cents",
            pay.get("amount_cents") == -420000, pay)
    je = next((x for x in rows if x.get("type") == "Journal Entry"), {})
    r.check("a row with no amount keeps null, never 0",
            "amount_cents" in je and je["amount_cents"] is None, je)
    r.check("window and as_of read the same way as the invoice list",
            (got.get("window_start"), got.get("window_end"), got.get("as_of"))
            == ("2026-01-01", "2026-03-31", "2026-04-07T09:15:00-04:00"),
            (got.get("window_start"), got.get("window_end"), got.get("as_of")))

    h2, b2 = _reorder(TXN_HEADER, TXN_BODY, [0, 9, 8, 7, 6, 5, 4, 3, 2, 1])
    txn2 = _xlsx(tmp / "txn2.xlsx", _titled(
        "Transaction List by Vendor", "January 1-March 31, 2026", h2, b2))
    err, g = _refusal(qx, qx.parse_vendor_transactions, txn2)
    r.check("vendor columns reordered -> the same rows",
            err is None and g.get("rows") == rows,
            str(err) if err else "rows differ: read by position")

    orphan = [[None, "01/08/2026", "No", "Bill", "9", "Yes", None, "AP",
               None, 1.0]] + TXN_BODY
    txn3 = _xlsx(tmp / "txn3.xlsx", _titled(
        "Transaction List by Vendor", "January 1-March 31, 2026", TXN_HEADER,
        orphan))
    err, g = _refusal(qx, qx.parse_vendor_transactions, txn3)
    r.check("a transaction above any vendor row is refused, not unattributed",
            err is not None, f"returned {g and len(g.get('rows') or [])} rows")

    sub = TXN_BODY[:5] + [["Subtotal freight", None, None, None, None, None,
                           None, None, None, 4200.0]] + TXN_BODY[5:]
    txn4 = _xlsx(tmp / "txn4.xlsx", _titled(
        "Transaction List by Vendor", "January 1-March 31, 2026", TXN_HEADER,
        sub))
    err, g = _refusal(qx, qx.parse_vendor_transactions, txn4)
    r.check("a summary row that is not 'Total for' is refused, not read as a "
            "vendor", err is not None,
            f"returned {g and len(g.get('rows') or [])} rows")

    # adversarial: a vendor whose NAME is "Total" is a vendor row (no date, no
    # amount), not a total -- skipping it files its bills under the vendor above
    named = TXN_BODY[:5] + [["Total", None, None, None, None, None, None, None,
                             None, None],
                            [None, "02/09/2026", "No", "Bill", "77", "Yes", None,
                             "Accounts Payable", "Freight", 10.0],
                            ["Total for Total", None, None, None, None, None,
                             None, None, None, 10.0]] + TXN_BODY[5:]
    txn5 = _xlsx(tmp / "txn5.xlsx", _titled(
        "Transaction List by Vendor", "January 1-March 31, 2026", TXN_HEADER,
        named))
    err, g = _refusal(qx, qx.parse_vendor_transactions, txn5)
    b77 = next((x for x in (g or {}).get("rows") or [] if x.get("num") == "77"),
               {})
    r.check("a vendor named 'Total' keeps its own bills",
            err is None and b77.get("vendor") == "Total",
            str(err) if err else b77)

    # adversarial: a Num cell Excel typed as a date comes back as a datetime,
    # which no JSON snapshot can hold -- refuse it at the parser, by name
    import datetime as _dt
    dated = [INV_BODY[0][:2] + [_dt.datetime(2026, 1, 5)] + INV_BODY[0][3:]]
    inv6 = _xlsx(tmp / "inv6.xlsx", _titled(
        "Invoice List by Date", "January 1-March 31, 2026", INV_HEADER, dated))
    err, g = _refusal(qx, qx.parse_invoice_list, inv6)
    r.check("a Num that is a date, not text, is refused", err is not None,
            f"returned {g!r}"[:160])

    err, _ = _refusal(qx, qx.parse_invoice_list, txn)
    r.check("a vendor list fed to the invoice parser is refused",
            err is not None and list(getattr(err, "found_header", None) or [])
            == TXN_HEADER, getattr(err, "found_header", None))


# -------------------------------------------------------------- snapshots --

def _snapshots(r, server, crm, tmp):
    r.section("snapshots live beside the store, replaced wholesale")
    save = getattr(server, "_save_qbo_snapshot", None)
    if not r.check("server has a snapshot writer", callable(save),
                   "no _save_qbo_snapshot"):
        return
    base = tmp / "machine"
    st = Store(server, base / "store")
    snapdir = base / "qbo-snapshots"
    store_before = sorted(p.name for p in st.path.iterdir())
    if not r.check("constructing the Store does not create the snapshot dir",
                   not snapdir.exists(), "Store.__init__ reached a sibling folder"):
        return

    rows1 = [{"type": "Invoice", "num": "7001", "amount_cents": 1250000,
              "open_cents": 0},
             {"type": "Invoice", "num": "7002", "amount_cents": 29,
              "open_cents": 29}]
    err = None
    try:
        path = save("invoices", "export", "2026-04-07T09:15:00-04:00",
                    "2026-01-01", "2026-03-31", rows1)
    except Exception as e:                            # noqa: BLE001
        err, path = f"{type(e).__name__}: {e}", None
    if not r.check("a snapshot saves", err is None, err or ""):
        return
    target = snapdir / "invoices.json"
    if not r.check("it is written to <store>/../qbo-snapshots/invoices.json",
                   target.exists() and Path(path).resolve() == target.resolve(),
                   f"wrote {path}"):
        return
    r.check("nothing new appeared inside the store",
            sorted(p.name for p in st.path.iterdir()) == store_before,
            sorted(p.name for p in st.path.iterdir()))
    doc = json.loads(target.read_text()) if target.exists() else {}
    r.check("the file carries kind, source, as_of, window, loaded_at, rows",
            {"kind", "source", "as_of", "window_start", "window_end",
             "loaded_at", "rows"} <= set(doc),
            sorted(doc))
    r.check("... with the values given",
            (doc.get("kind"), doc.get("source"), doc.get("as_of"),
             doc.get("window_start"), doc.get("window_end"), doc.get("rows"))
            == ("invoices", "export", "2026-04-07T09:15:00-04:00",
                "2026-01-01", "2026-03-31", rows1), doc)
    r.check("loaded_at is a timestamp", isinstance(doc.get("loaded_at"), str)
            and doc["loaded_at"][:2] == "20", doc.get("loaded_at"))

    rows2 = rows1[:1]
    save("invoices", "connector", "2026-04-08", "2026-01-01", "2026-04-08",
         rows2)
    doc = json.loads(target.read_text())
    r.check("a second load replaces the file wholesale, never appends",
            doc.get("rows") == rows2 and doc.get("source") == "connector",
            doc.get("rows"))
    r.check("no temp file is left beside it",
            sorted(p.name for p in snapdir.iterdir()) == ["invoices.json"],
            sorted(p.name for p in snapdir.iterdir()))

    r.section("the writer refuses what it cannot store honestly")
    for label, args in (
            ("an unknown kind", ("payments", "export", "2026-04-07", "2026-01-01",
                                 "2026-03-31", rows1)),
            ("an unknown source", ("invoices", "email", "2026-04-07",
                                   "2026-01-01", "2026-03-31", rows1)),
            ("float cents", ("invoices", "export", "2026-04-07", "2026-01-01",
                             "2026-03-31",
                             [{"type": "Invoice", "num": "1",
                               "amount_cents": 12.5}])),
            ("a boolean in a cents field", ("invoices", "export", "2026-04-07",
                                            "2026-01-01", "2026-03-31",
                                            [{"type": "Invoice", "num": "1",
                                              "amount_cents": True}])),
            ("a value JSON cannot hold (a date object)",
             ("invoices", "export", "2026-04-07", "2026-01-01", "2026-03-31",
              [{"type": "Invoice", "num": __import__("datetime").date(2026, 1, 5),
                "amount_cents": 100}]))):
        prior = target.read_bytes()
        try:
            save(*args)
            refused = False
        except server.StoreError:
            refused = True
        except Exception:                             # noqa: BLE001
            refused = False                           # a crash is not a refusal
        r.check(f"{label} is refused with a StoreError", refused)
        r.check(f"... and the snapshot on disk is untouched ({label})",
                target.read_bytes() == prior)

    # adversarial: the snapshot path is occupied by a plain FILE -- the writer
    # must refuse in a sentence (StoreError), not raise a raw OSError
    other = tmp / "machine2"
    st2 = Store(server, other / "store")
    (other / "qbo-snapshots").write_text("not a folder")
    try:
        save("invoices", "export", "2026-04-07", "2026-01-01", "2026-03-31",
             rows1)
        outcome = "saved"
    except server.StoreError:
        outcome = "StoreError"
    except Exception as e:                            # noqa: BLE001
        outcome = type(e).__name__
    r.check("a file where the snapshot folder should be -> StoreError",
            outcome == "StoreError", outcome)
    r.check("... and that file is left alone",
            (other / "qbo-snapshots").read_text() == "not a folder")
    st.rebind()

    r.section("the store's Windows-lock retry, not a bare os.replace")
    real = os.replace
    calls = {"n": 0}

    def flaky(a, b):
        if str(b).endswith("invoices.json") and calls["n"] < 2:
            calls["n"] += 1
            raise PermissionError("locked by OneDrive")
        return real(a, b)

    os.replace = flaky
    try:
        err = None
        try:
            save("invoices", "export", "2026-04-09", "2026-01-01", "2026-04-09",
                 rows1)
        except Exception as e:                        # noqa: BLE001
            err = f"{type(e).__name__}: {e}"
    finally:
        os.replace = real
    r.check("two PermissionErrors then success -> saved",
            err is None and json.loads(target.read_text()).get("as_of")
            == "2026-04-09", err or "")

    prior = target.read_bytes()

    def locked(a, b):
        if str(b).endswith("invoices.json"):
            raise PermissionError("locked by OneDrive")
        return real(a, b)

    os.replace = locked
    try:
        try:
            save("invoices", "export", "2026-04-10", "2026-01-01",
                 "2026-04-10", rows2)
            outcome = "saved"
        except server.StoreError:
            outcome = "StoreError"
        except Exception as e:                        # noqa: BLE001
            outcome = f"{type(e).__name__}"
    finally:
        os.replace = real
    r.check("a lock that never clears -> StoreError, not a raw PermissionError",
            outcome == "StoreError", outcome)
    r.check("... the prior snapshot survives and no temp is left",
            target.read_bytes() == prior
            and sorted(p.name for p in snapdir.iterdir()) == ["invoices.json"],
            sorted(p.name for p in snapdir.iterdir()))

    r.section("the store and the importer never touch the snapshot dir")
    before = {p.name: p.read_bytes() for p in snapdir.iterdir()}
    st.rebind()                                       # a relaunch
    r.check("a Store relaunch leaves the snapshot dir byte-identical",
            {p.name: p.read_bytes() for p in snapdir.iterdir()} == before)
    nrm = _load(crm, "normalize")
    if nrm is not None:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sales Tracker 2026"
        ws.append(["Project#", "Date", "Customer", "Description", "Location",
                   "Status", "PO Y/N", "Client PO#", "Invoice #", "Revenue",
                   "Total Cost", "Total GP", "Margain", "Notes"])
        ws.append(["9001", None, "Brightwater Fabrication", "Line install",
                   "Dayton OH", "won", "Y", "PO-5501", "8801",
                   50000, 30000, 20000, 0.4, ""])
        pt = wb.create_sheet("Project Tracker")
        pt.append(["Unrivaled Project#:", "Client PO#:", "Start Date:",
                   "Client Name:", "Client Location:", "Open Orders Notes:",
                   "Vendor 1 PO#:", "Vendor 1 Ship Date:"])
        pt.append(["9001", "PO-5501", None, "Brightwater Fabrication",
                   "Dayton OH", "waiting on frames", "VPO-1", None])
        wb.create_sheet("Client Contacts").append(
            ["Client Business", "Client Name", "Email", "Phone Number",
             "Job Title", "Location", "Action Taken and Notes",
             "Last Date of Action"])
        wb.create_sheet("Vendor Contacts").append(
            ["Company", "Headquarters Location", "Sales Rep/Contact",
             "Contact Email", "Contact Phone Number", "Offerings",
             "Send PO's to", "Send Invoices to"])
        x = tmp / "tracker.xlsx"
        wb.save(x)
        pdir = str(Path(crm) / "pipeline")
        sys.path.insert(0, pdir)
        err = None
        try:
            nrm.run(str(x), str(st.path), force=True, mode="merge")
        except BaseException as e:                    # noqa: BLE001
            err = f"{type(e).__name__}: {e}"
        finally:
            sys.path.remove(pdir)
        wrote = err is None and any(
            str(p.get("project_no")) == "9001" for p in st.read("projects") or [])
        r.check("an import into the store runs and writes the store", wrote,
                err or "project 9001 not in projects.json")
        r.check("the import (normalize + merge) leaves the snapshot dir "
                "byte-identical",
                wrote and {p.name: p.read_bytes() for p in snapdir.iterdir()} == before,
                sorted(p.name for p in snapdir.iterdir()))
